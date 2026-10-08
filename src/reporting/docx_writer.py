"""Word (.docx) output of a built Report (python-docx, A4 portrait)."""

import io
from typing import List

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .builder import Figure, KeyValues, Paragraph, Report, Table

ACCENT = RGBColor(0x1F, 0x6B, 0x73)
ACCENT_HEX = "1F6B73"
ROW_ALT_HEX = "F4F4F2"
BORDER_HEX = "D9D8D2"
INK = RGBColor(0x0B, 0x0B, 0x0B)
INK_SECONDARY = RGBColor(0x52, 0x51, 0x4E)
CONTENT_WIDTH = Cm(17.0)


def write_docx(report: Report, path: str):
    t = report.texts
    doc = Document()
    _page_setup(doc)
    _styles(doc)
    _footer(doc, report)

    _cover(doc, report)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    doc.add_paragraph(t["contents"], style="TOC Heading")
    _toc_field(doc, [section.title for section in report.sections], t["toc_hint"])
    _update_fields_on_open(doc)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    counters = {"figure": 0, "table": 0}
    for index, section in enumerate(report.sections):
        if index > 0:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        doc.add_heading(section.title, level=1)
        for block in section.blocks:
            if isinstance(block, Paragraph):
                doc.add_paragraph(block.text)
            elif isinstance(block, KeyValues):
                _key_values(doc, block.rows)
            elif isinstance(block, Table):
                counters["table"] += 1
                _caption(doc, f"{t['table']} {counters['table']}: {block.caption}", keep_with_next=True)
                _table(doc, block.headers, block.rows)
                doc.add_paragraph()
            elif isinstance(block, Figure):
                counters["figure"] += 1
                picture = doc.add_paragraph()
                picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
                picture.paragraph_format.keep_with_next = True
                picture.add_run().add_picture(io.BytesIO(block.png), width=CONTENT_WIDTH)
                _caption(doc, f"{t['figure']} {counters['figure']}: {block.caption}")

    doc.save(path)


# ----------------------------------------------------------------------
def _page_setup(doc):
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2.0)
    section.top_margin = section.bottom_margin = Cm(2.0)


def _styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Calibri")
    normal.paragraph_format.space_after = Pt(6)
    for name, size in (("Heading 1", 18), ("Heading 2", 14)):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = ACCENT
        style.font.bold = True
    caption = doc.styles["Caption"]
    caption.font.size = Pt(9.5)
    caption.font.italic = False
    caption.font.color.rgb = INK_SECONDARY


def _cover(doc, report: Report):
    t = report.texts
    for _ in range(6):
        doc.add_paragraph()
    title = doc.add_paragraph()
    run = title.add_run(report.title)
    run.font.size = Pt(30)
    run.font.bold = True
    run.font.color.rgb = INK
    rule = doc.add_paragraph()
    _bottom_border(rule, ACCENT_HEX, 18)
    rows = []
    if report.project:
        rows.append((t["project"], report.project))
    if report.author:
        rows.append((t["prepared_by"], report.author))
    rows.append((t["date"], report.date))
    if report.file_name:
        rows.append((t["data_file"], report.file_name))
    rows.append((t["range"], report.range_text))
    for key, value in rows:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        k = p.add_run(f"{key}: ")
        k.font.size = Pt(12)
        k.font.color.rgb = INK_SECONDARY
        v = p.add_run(value)
        v.font.size = Pt(12)


def _footer(doc, report: Report):
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    label = footer.add_run(f"{report.title}  ·  {report.texts['page']} ")
    label.font.size = Pt(9)
    label.font.color.rgb = INK_SECONDARY
    _field(footer, "PAGE", shown="1", size=9, color=INK_SECONDARY)


def _field(paragraph, instruction: str, shown: str = "", size=None, color=None):
    """
    A Word field (PAGE, TOC, ...) that Word computes; `shown` is displayed
    until it does (and by viewers that do not compute fields).
    """
    def run_with(*elements):
        run = paragraph.add_run()
        if size:
            run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color
        for element in elements:
            run._r.append(element)
        return run

    def char(kind):
        element = OxmlElement("w:fldChar")
        element.set(qn("w:fldCharType"), kind)
        return element

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    run_with(char("begin"))
    run_with(instr)
    run_with(char("separate"))
    shown_run = run_with()
    shown_run.text = shown
    run_with(char("end"))


def _toc_field(doc, entries: List[str], hint: str):
    """Table of contents field, showing the section titles until Word updates it."""
    paragraph = doc.add_paragraph()
    _field(paragraph, r'TOC \o "1-2" \h \z \u', shown="\n".join(entries), size=11)
    note = doc.add_paragraph()
    run = note.add_run(hint)
    run.font.size = Pt(8.5)
    run.font.color.rgb = INK_SECONDARY


def _update_fields_on_open(doc):
    """Ask Word to refresh fields (table of contents, page numbers) on open."""
    settings = doc.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)


def _caption(doc, text: str, keep_with_next: bool = False):
    p = doc.add_paragraph(text, style="Caption")
    p.paragraph_format.keep_with_next = keep_with_next
    if not keep_with_next:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def _shade(cell, hex_color: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), hex_color)
    tc_pr.append(shading)


def _bottom_border(paragraph, hex_color: str, size: int):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", str(size)), ("w:space", "1"), ("w:color", hex_color)):
        bottom.set(qn(k), v)
    borders.append(bottom)
    p_pr.append(borders)


def _borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "bottom", "insideH"):
        element = OxmlElement(f"w:{edge}")
        for k, v in (("w:val", "single"), ("w:sz", "4"), ("w:space", "0"), ("w:color", BORDER_HEX)):
            element.set(qn(k), v)
        borders.append(element)
    for edge in ("left", "right", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "nil")
        borders.append(element)
    tbl_pr.append(borders)


def _write_cell(cell, text: str, bold=False, color=INK, align=WD_ALIGN_PARAGRAPH.LEFT, size=9.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def _table(doc, headers: List[str], rows: List[List[str]]):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _borders(table)
    size = 9.5 if len(headers) <= 7 else 8.5
    header_row = table.rows[0]
    _repeat_header(header_row)
    for c, header in enumerate(headers):
        cell = header_row.cells[c]
        _write_cell(cell, header, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF),
                    align=WD_ALIGN_PARAGRAPH.LEFT if c == 0 else WD_ALIGN_PARAGRAPH.RIGHT, size=size)
        _shade(cell, ACCENT_HEX)
    for r, row in enumerate(rows, start=1):
        for c, text in enumerate(row):
            cell = table.rows[r].cells[c]
            _write_cell(cell, text, align=WD_ALIGN_PARAGRAPH.LEFT if c == 0 else WD_ALIGN_PARAGRAPH.RIGHT,
                        size=size)
            if r % 2 == 0:
                _shade(cell, ROW_ALT_HEX)
    first = Cm(5.0 if len(headers) > 4 else 7.0)
    other = int((CONTENT_WIDTH - first) / max(len(headers) - 1, 1))
    for row in table.rows:
        row.cells[0].width = first
        for cell in row.cells[1:]:
            cell.width = other


def _repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def _key_values(doc, rows):
    table = doc.add_table(rows=len(rows), cols=2)
    _borders(table)
    for r, (key, value) in enumerate(rows):
        _write_cell(table.rows[r].cells[0], key, bold=True, color=INK_SECONDARY, size=10.5)
        _write_cell(table.rows[r].cells[1], value, size=10.5)
        table.rows[r].cells[0].width = Cm(6.0)
        table.rows[r].cells[1].width = CONTENT_WIDTH - Cm(6.0)
    doc.add_paragraph()
