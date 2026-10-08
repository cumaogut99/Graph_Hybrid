"""PowerPoint (.pptx) output of a built Report (python-pptx, 16:9)."""

import io
from typing import List, Tuple

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

from .builder import Figure, KeyValues, Paragraph, Report, Table

ACCENT = RGBColor(0x1F, 0x6B, 0x73)       # app accent (teal)
ACCENT_LIGHT = RGBColor(0xE6, 0xF1, 0xF2)
INK = RGBColor(0x0B, 0x0B, 0x0B)
INK_SECONDARY = RGBColor(0x52, 0x51, 0x4E)
INK_MUTED = RGBColor(0x89, 0x87, 0x81)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ROW_ALT = RGBColor(0xF4, 0xF4, 0xF2)

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.6)
CONTENT_TOP = Inches(1.55)
CONTENT_BOTTOM = Inches(6.85)
TABLE_ROWS_PER_SLIDE = 11
NOTE_LINE = Inches(0.3)

# A planned slide: [kind, section title, payload, notes]; notes are short
# texts shown at the bottom of the slide instead of on a slide of their own
Slide = list


def write_pptx(report: Report, path: str):
    t = report.texts
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    blank = prs.slide_layouts[6]

    plan = _plan(report)
    # Slide numbers of the sections (1: cover, 2: contents)
    first_slide = {}
    for number, (_, section_title, _, _) in enumerate(plan, start=3):
        first_slide.setdefault(section_title, number)

    _cover(prs.slides.add_slide(blank), report)
    contents = prs.slides.add_slide(blank)
    _header(contents, t["contents"], "")
    _contents(contents, [(s.title, first_slide.get(s.title)) for s in report.sections if s.title in first_slide])
    _footer(contents, report, 2)

    for number, (kind, section_title, payload, notes) in enumerate(plan, start=3):
        slide = prs.slides.add_slide(blank)
        notes_height = NOTE_LINE * len(notes)
        if kind == "summary":
            key_values, paragraphs = payload
            _header(slide, section_title, "")
            _summary(slide, key_values, paragraphs)
        elif kind == "table":
            table, rows, part, parts = payload
            caption = table.caption + (f" ({part}/{parts})" if parts > 1 else "")
            _header(slide, section_title, caption)
            _table(slide, table.headers, rows)
        elif kind == "figure":
            _header(slide, section_title, payload.caption)
            _picture(slide, payload.png, CONTENT_BOTTOM - notes_height)
        elif kind == "text":
            _header(slide, section_title, "")
            _text(slide, payload)
        if notes:
            _notes(slide, notes)
        _footer(slide, report, number)

    prs.save(path)


# ----------------------------------------------------------------------
# Planning: content blocks -> slides
# ----------------------------------------------------------------------
def _plan(report: Report) -> List[Slide]:
    slides: List[Slide] = []
    for section in report.sections:
        first_of_section = len(slides)
        key_values = [b for b in section.blocks if isinstance(b, KeyValues)]
        if key_values:
            # Summary-like section: facts and its texts on one slide
            paragraphs = [b.text for b in section.blocks if isinstance(b, Paragraph)]
            slides.append(["summary", section.title, (key_values[0], paragraphs), []])
            others = [b for b in section.blocks if isinstance(b, (Table, Figure))]
        else:
            others = section.blocks

        pending: List[str] = []
        for block in others:
            if isinstance(block, Paragraph):
                pending.append(block.text)
                continue
            first_of_block = len(slides)
            if isinstance(block, Table):
                chunks = [block.rows[i:i + TABLE_ROWS_PER_SLIDE]
                          for i in range(0, len(block.rows), TABLE_ROWS_PER_SLIDE)] or [[]]
                for part, rows in enumerate(chunks, 1):
                    slides.append(["table", section.title, (block, rows, part, len(chunks)), []])
            elif isinstance(block, Figure):
                slides.append(["figure", section.title, block, []])
            slides[first_of_block][3] = pending  # text before a block: note on its slide
            pending = []
        if pending:
            if len(slides) > first_of_section:
                slides[-1][3] = slides[-1][3] + pending  # closing text: note on the last slide
            else:
                slides.append(["text", section.title, pending, []])
    return slides


# ----------------------------------------------------------------------
# Slide parts
# ----------------------------------------------------------------------
def _textbox(slide, left, top, width, height, text, size, color=INK, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    frame.margin_left = frame.margin_right = 0
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def _rect(slide, left, top, width, height, color):
    from pptx.enum.shapes import MSO_SHAPE
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def _cover(slide, report: Report):
    t = report.texts
    _rect(slide, 0, 0, Inches(0.35), SLIDE_H, ACCENT)
    _textbox(slide, Inches(1.2), Inches(2.0), Inches(11), Inches(1.4), report.title, 40, INK, bold=True,
             anchor=MSO_ANCHOR.BOTTOM)
    _rect(slide, Inches(1.2), Inches(3.55), Inches(1.6), Inches(0.06), ACCENT)
    lines = []
    if report.project:
        lines.append(f"{t['project']}: {report.project}")
    if report.author:
        lines.append(f"{t['prepared_by']}: {report.author}")
    lines.append(f"{t['date']}: {report.date}")
    if report.file_name:
        lines.append(f"{t['data_file']}: {report.file_name}")
    lines.append(f"{t['range']}: {report.range_text}")
    box = _textbox(slide, Inches(1.2), Inches(3.85), Inches(11), Inches(2.6), lines[0], 18, INK_SECONDARY)
    for line in lines[1:]:
        p = box.text_frame.add_paragraph()
        run = p.add_run()
        run.text = line
        run.font.size = Pt(16)
        run.font.color.rgb = INK_SECONDARY
        p.space_before = Pt(6)


def _header(slide, title: str, subtitle: str):
    _textbox(slide, MARGIN, Inches(0.35), SLIDE_W - 2 * MARGIN, Inches(0.6), title, 26, INK, bold=True)
    _rect(slide, MARGIN, Inches(0.98), Inches(0.9), Inches(0.05), ACCENT)
    if subtitle:
        _textbox(slide, MARGIN, Inches(1.08), SLIDE_W - 2 * MARGIN, Inches(0.42), subtitle, 14, INK_SECONDARY)


def _footer(slide, report: Report, number: int):
    footer = report.title + (f"  ·  {report.project}" if report.project else "")
    _textbox(slide, MARGIN, Inches(7.0), Inches(10), Inches(0.3), footer, 10, INK_MUTED)
    _textbox(slide, SLIDE_W - MARGIN - Inches(1.5), Inches(7.0), Inches(1.5), Inches(0.3),
             str(number), 10, INK_MUTED, align=PP_ALIGN.RIGHT)


def _contents(slide, entries: List[Tuple[str, int]]):
    top = CONTENT_TOP
    for title, number in entries:
        _textbox(slide, MARGIN, top, Inches(9), Inches(0.5), title, 20, INK)
        _textbox(slide, SLIDE_W - MARGIN - Inches(1.5), top, Inches(1.5), Inches(0.5),
                 str(number), 20, ACCENT, bold=True, align=PP_ALIGN.RIGHT)
        top += Inches(0.65)


def _summary(slide, key_values: KeyValues, paragraphs: List[str]):
    width = Inches(7.4) if paragraphs else SLIDE_W - 2 * MARGIN
    _key_value_table(slide, key_values.rows, MARGIN, CONTENT_TOP, width)
    if paragraphs:
        left = MARGIN + width + Inches(0.4)
        _text(slide, paragraphs, left=left, width=SLIDE_W - MARGIN - left)


def _key_value_table(slide, rows, left, top, width):
    shape = slide.shapes.add_table(len(rows), 2, left, top, width, Inches(0.42) * len(rows))
    table = shape.table
    table.first_row = False
    table.columns[0].width = int(width * 0.38)
    table.columns[1].width = width - table.columns[0].width
    for r, (key, value) in enumerate(rows):
        for c, text in enumerate((key, value)):
            cell = table.cell(r, c)
            _cell(cell, text, 14, INK_SECONDARY if c == 0 else INK, bold=(c == 0),
                  fill=ROW_ALT if r % 2 == 0 else WHITE)


def _text(slide, paragraphs: List[str], left=MARGIN, width=None):
    width = width or SLIDE_W - 2 * MARGIN
    box = _textbox(slide, left, CONTENT_TOP, width, CONTENT_BOTTOM - CONTENT_TOP, paragraphs[0], 16, INK)
    for text in paragraphs[1:]:
        p = box.text_frame.add_paragraph()
        p.space_before = Pt(12)
        run = p.add_run()
        run.text = text
        run.font.size = Pt(16)
        run.font.color.rgb = INK


def _cell(cell, text, size, color, bold=False, fill=None, align=PP_ALIGN.LEFT):
    cell.text = ""
    frame = cell.text_frame
    frame.word_wrap = True
    p = frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = cell.margin_right = Inches(0.08)
    cell.margin_top = cell.margin_bottom = Inches(0.03)
    if fill is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill


def _table(slide, headers: List[str], rows: List[List[str]]):
    width = SLIDE_W - 2 * MARGIN
    n_cols = len(headers)
    size = 13 if n_cols <= 7 else 11
    shape = slide.shapes.add_table(len(rows) + 1, n_cols, MARGIN, CONTENT_TOP, width,
                                   Inches(0.4) * (len(rows) + 1))
    table = shape.table
    first = int(width * (0.28 if n_cols > 4 else 0.4))
    table.columns[0].width = first
    for c in range(1, n_cols):
        table.columns[c].width = int((width - first) / (n_cols - 1))
    for c, header in enumerate(headers):
        _cell(table.cell(0, c), header, size, WHITE, bold=True, fill=ACCENT,
              align=PP_ALIGN.LEFT if c == 0 else PP_ALIGN.RIGHT)
    for r, row in enumerate(rows, start=1):
        for c, text in enumerate(row):
            _cell(table.cell(r, c), text, size, INK, fill=ROW_ALT if r % 2 == 0 else WHITE,
                  align=PP_ALIGN.LEFT if c == 0 else PP_ALIGN.RIGHT)


def _notes(slide, notes: List[str]):
    top = CONTENT_BOTTOM - NOTE_LINE * len(notes) + Inches(0.05)
    box = _textbox(slide, MARGIN, top, SLIDE_W - 2 * MARGIN, NOTE_LINE * len(notes), notes[0], 12, INK_SECONDARY)
    for text in notes[1:]:
        run = box.text_frame.add_paragraph().add_run()
        run.text = text
        run.font.size = Pt(12)
        run.font.color.rgb = INK_SECONDARY


def _picture(slide, png: bytes, bottom=CONTENT_BOTTOM):
    with Image.open(io.BytesIO(png)) as image:
        w_px, h_px = image.size
    max_w = SLIDE_W - 2 * MARGIN
    max_h = bottom - CONTENT_TOP
    scale = min(max_w / w_px, max_h / h_px)
    width, height = int(w_px * scale), int(h_px * scale)
    left = int((SLIDE_W - width) / 2)
    slide.shapes.add_picture(io.BytesIO(png), Emu(left), CONTENT_TOP, Emu(width), Emu(height))
