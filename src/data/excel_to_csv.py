#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Excel → CSV Dönüştürücü
=======================

Excel dosyaları (.xlsx, .xlsm, .xls, .ods) uygulamanın CSV hattından geçer:
seçilen sayfa olduğu gibi (header yorumu yapılmadan, açıklama satırları
dahil) geçici bir UTF-8 CSV'ye yazılır. Böylece import dialog'undaki
format algılama, header/veri satırı seçimi ve CSV → MPAI dönüştürücü
Excel için de aynen çalışır.

Okuma fastexcel (calamine) ile yapılır: hızlıdır ve .xls dahil tüm
formatları okur. fastexcel yoksa .xlsx/.xlsm için openpyxl kullanılır.
"""

import csv
import datetime
import logging
import os
import shutil
import tempfile
from typing import List, Optional, Union

logger = logging.getLogger(__name__)

EXCEL_EXTENSIONS = ('.xlsx', '.xlsm', '.xls', '.ods')

try:
    import fastexcel
    HAS_FASTEXCEL = True
except ImportError:
    HAS_FASTEXCEL = False


def is_excel_file(file_path: str) -> bool:
    return os.path.splitext(file_path)[1].lower() in EXCEL_EXTENSIONS


def list_sheets(file_path: str) -> List[str]:
    """Excel dosyasındaki sayfa adlarını döndür."""
    if HAS_FASTEXCEL:
        return list(fastexcel.read_excel(file_path).sheet_names)
    wb = _open_openpyxl(file_path)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def excel_to_csv(file_path: str, csv_path: str, sheet: Union[str, int, None] = None) -> int:
    """
    Bir Excel sayfasını ham haliyle CSV'ye yaz (virgül ayırıcı, nokta ondalık).

    Args:
        file_path: Excel dosyası
        csv_path: Yazılacak CSV yolu
        sheet: Sayfa adı veya indeksi (None = ilk sayfa)

    Returns:
        Yazılan satır sayısı
    """
    sheet = 0 if sheet in (None, '') else sheet
    if HAS_FASTEXCEL:
        rows = _write_with_fastexcel(file_path, csv_path, sheet)
    else:
        rows = _write_with_openpyxl(file_path, csv_path, sheet)
    logger.info(f"[EXCEL] {os.path.basename(file_path)} [{sheet}] -> {csv_path} ({rows} satır)")
    return rows


def excel_to_temp_csv(file_path: str, sheet: Union[str, int, None] = None) -> str:
    """Sayfayı yeni bir geçici klasördeki CSV'ye yaz ve yolunu döndür."""
    temp_dir = tempfile.mkdtemp(prefix="timegraph_excel_")
    csv_path = os.path.join(temp_dir, os.path.splitext(os.path.basename(file_path))[0] + ".csv")
    try:
        excel_to_csv(file_path, csv_path, sheet)
    except Exception:
        remove_temp_csv(csv_path)
        raise
    return csv_path


def remove_temp_csv(csv_path: Optional[str]):
    """excel_to_temp_csv ile oluşturulan CSV'yi ve klasörünü sil."""
    if not csv_path:
        return
    temp_dir = os.path.dirname(csv_path)
    try:
        if os.path.basename(temp_dir).startswith("timegraph_excel_"):
            shutil.rmtree(temp_dir, ignore_errors=True)
        elif os.path.exists(csv_path):
            os.remove(csv_path)
    except Exception as e:
        logger.debug(f"[EXCEL] Temp CSV silinemedi: {e}")


def read_sheet_rows(file_path: str, sheet: Union[str, int]) -> List[List[str]]:
    """Bir Excel sayfasının tüm satırlarını metin olarak döndür (başlık ayrımı yapmadan)."""
    if HAS_FASTEXCEL:
        df = fastexcel.read_excel(file_path).load_sheet(
            sheet, header_row=None, dtypes="string"
        ).to_polars()
        return [['' if v is None else v for v in row] for row in df.iter_rows()]
    wb = _open_openpyxl(file_path)
    try:
        ws = wb.worksheets[sheet] if isinstance(sheet, int) else wb[sheet]
        return [['' if v is None else _format_cell(v) for v in row]
                for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def _write_with_fastexcel(file_path: str, csv_path: str, sheet: Union[str, int]) -> int:
    # header_row=None: tüm satırlar veri olarak gelir (açıklama satırları dahil)
    # dtypes="string": karışık kolonlar (metin + sayı) kaybolmasın
    excel_sheet = fastexcel.read_excel(file_path).load_sheet(
        sheet, header_row=None, dtypes="string"
    )
    df = excel_sheet.to_polars()
    df.write_csv(csv_path, include_header=False)
    return df.height


def _open_openpyxl(file_path: str):
    ext = os.path.splitext(file_path)[1].lower()
    if ext not in ('.xlsx', '.xlsm'):
        raise ImportError(
            f"{ext} dosyaları için fastexcel gereklidir.\nKurulum: pip install fastexcel"
        )
    import openpyxl
    return openpyxl.load_workbook(file_path, read_only=True, data_only=True)


def _write_with_openpyxl(file_path: str, csv_path: str, sheet: Union[str, int]) -> int:
    wb = _open_openpyxl(file_path)
    try:
        ws = wb.worksheets[sheet] if isinstance(sheet, int) else wb[sheet]
        rows = 0
        with open(csv_path, 'w', encoding='utf-8', newline='') as f:
            writer = csv.writer(f)
            for row in ws.iter_rows(values_only=True):
                writer.writerow(['' if v is None else _format_cell(v) for v in row])
                rows += 1
        return rows
    finally:
        wb.close()


def _format_cell(value) -> str:
    if isinstance(value, datetime.datetime):
        return value.isoformat(sep=' ')
    if isinstance(value, (datetime.date, datetime.time)):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)
